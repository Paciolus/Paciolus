"""
Sprint 813 — rendering-equivalence harness for every memo PDF.

Sprint 813's structural items (``render_memo``, ``ledger_table_style``
params, ``source_line``, the reference-prefix argument, the cached styles)
are refactors: the PDFs a customer downloads must not change. This test
pins that. Each case below renders one memo from a fixed fixture, with the
clock frozen, and compares what came out against a snapshot recorded on the
pre-refactor code (``fixtures/memo_render_snapshot.json``):

* page count, and each page's media box;
* each page's extracted text;
* each text run's position, font and size (a digest) plus the bounding box
  of all text — a margin or leading change moves these;
* each page's content stream (a digest) — rules, fills and table grids are
  drawn there, so a table-style change shows up even when no text moves.

It never compares PDF bytes, which carry a creation date and document ID.

Every case renders twice: with no branding, and under an Enterprise
branding context (header, footer and logo), because the page chrome is
exactly what ``render_memo`` takes over.

To re-record after an *intended* output change:
    MEMO_SNAPSHOT_UPDATE=1 pytest tests/test_memo_render_equivalence.py
and review the snapshot diff like any other output change.
"""

import hashlib
import io
import json
import os
import sys
import zlib
from collections.abc import Callable
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

import pytest
from freezegun import freeze_time
from pypdf import PdfReader

sys.path.insert(0, str(Path(__file__).parent.parent))

from shared.pdf_branding import PDFBrandingContext, apply_pdf_branding

SNAPSHOT_PATH = Path(__file__).parent / "fixtures" / "memo_render_snapshot.json"
UPDATE = os.environ.get("MEMO_SNAPSHOT_UPDATE") == "1"
FROZEN_AT = "2026-03-02 09:30:15"


def _tiny_png() -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (40, 12), (33, 33, 33)).save(buf, format="PNG")
    return buf.getvalue()


BRANDINGS: dict[str, PDFBrandingContext] = {
    "default": PDFBrandingContext(),
    "branded": PDFBrandingContext(
        tier_has_branding=True,
        logo_bytes=_tiny_png(),
        logo_content_type="image/png",
        header_text="ACME Audit Partners LLP",
        footer_text="Confidential - Prepared for ACME Client Corp.",
    ),
}


# ---------------------------------------------------------------------------
# Measurement
# ---------------------------------------------------------------------------


def _digest(payload: Any) -> str:
    return hashlib.sha256(json.dumps(payload, sort_keys=True).encode()).hexdigest()[:16]


def _measure(pdf: bytes) -> dict[str, Any]:
    reader = PdfReader(io.BytesIO(pdf))
    pages = []
    for page in reader.pages:
        runs: list[list] = []

        def visit(text, cm, tm, font_dict, font_size, runs=runs):
            if not text.strip():
                return
            x = cm[0] * tm[4] + cm[2] * tm[5] + cm[4]
            y = cm[1] * tm[4] + cm[3] * tm[5] + cm[5]
            font = (font_dict or {}).get("/BaseFont", "")
            runs.append([round(x, 2), round(y, 2), str(font), round(float(font_size), 2), text])

        text = page.extract_text(visitor_text=visit) or ""
        contents = page.get_contents()
        stream = contents.get_data() if contents is not None else b""
        xs = [r[0] for r in runs] or [0]
        ys = [r[1] for r in runs] or [0]
        pages.append(
            {
                "size": [round(float(v), 2) for v in page.mediabox],
                "text": text,
                "text_bbox": [min(xs), min(ys), max(xs), max(ys)],
                "runs": _digest(runs),
                "content": hashlib.sha256(stream).hexdigest()[:16],
            }
        )
    return {"page_count": len(pages), "pages": pages}


# ---------------------------------------------------------------------------
# Cases — every memo module, plus the anomaly summary and TB/FS reports
# ---------------------------------------------------------------------------


def _sample(generator_name: str) -> Callable[[Any], bytes]:
    """Render one ``generate_sample_reports`` generator and return its PDF."""

    def render(_request: Any) -> bytes:
        import generate_sample_reports as samples

        captured: list[bytes] = []
        original = samples.save_pdf
        samples.save_pdf = lambda _name, pdf: captured.append(pdf)
        try:
            getattr(samples, generator_name)()
        finally:
            samples.save_pdf = original
        assert len(captured) == 1, generator_name
        return captured[0]

    return render


SAMPLE_GENERATORS = [
    "gen_tb_diagnostic",
    "gen_financial_statements",
    "gen_je_testing",
    "gen_ap_testing",
    "gen_payroll_testing",
    "gen_revenue_testing",
    "gen_ar_aging",
    "gen_fixed_asset_testing",
    "gen_inventory_testing",
    "gen_bank_rec",
    "gen_three_way_match",
    "gen_multi_period",
    "gen_sampling_design",
    "gen_sampling_evaluation",
    "gen_currency_conversion",
    "gen_preflight",
    "gen_population_profile",
    "gen_expense_category",
    "gen_accrual_completeness",
    "gen_flux_expectations",
    "gen_anomaly_summary",
]


def _engagement_dashboard(_request: Any) -> bytes:
    from engagement_dashboard_memo import generate_engagement_dashboard_memo
    from tests.test_engagement_dashboard_memo import _make_dashboard_result

    return generate_engagement_dashboard_memo(
        _make_dashboard_result(),
        client_name="Meridian Capital Group",
        period_tested="FY 2025",
        prepared_by="John Doe",
        reviewed_by="Jane Smith",
        workpaper_date="2025-12-31",
        include_signoff=True,
    )


def _engagement(request: Any):
    from engagement_manager import EngagementManager

    db = request.getfixturevalue("db_session")
    user = request.getfixturevalue("make_user")()
    client = request.getfixturevalue("make_client")(user=user, name="Acme Corp")
    eng = EngagementManager(db).create_engagement(
        user.id,
        client.id,
        datetime(2025, 1, 1, tzinfo=UTC),
        datetime(2025, 12, 31, tzinfo=UTC),
        materiality_amount=100000.0,
    )
    return db, user, eng


def _analytical_expectations(request: Any) -> bytes:
    from analytical_expectation_memo_generator import AnalyticalExpectationMemoGenerator
    from analytical_expectations_manager import AnalyticalExpectationsManager
    from analytical_expectations_model import ExpectationTargetType

    db, user, eng = _engagement(request)
    mgr = AnalyticalExpectationsManager(db)
    mgr.create_expectation(
        user_id=user.id,
        engagement_id=eng.id,
        procedure_target_type=ExpectationTargetType.ACCOUNT,
        procedure_target_label="Revenue",
        expected_value=1000.0,
        precision_threshold_amount=50.0,
        corroboration_basis_text="Prior-period growth",
        corroboration_tags=["prior_period"],
    )
    within = mgr.create_expectation(
        user_id=user.id,
        engagement_id=eng.id,
        procedure_target_type=ExpectationTargetType.RATIO,
        procedure_target_label="Current ratio",
        expected_value=2.0,
        precision_threshold_percent=10.0,
        corroboration_basis_text="Industry median",
        corroboration_tags=["industry_data", "prior_period"],
    )
    mgr.update_expectation(user_id=user.id, expectation_id=within.id, result_actual_value=2.05)
    exceeds = mgr.create_expectation(
        user_id=user.id,
        engagement_id=eng.id,
        procedure_target_type=ExpectationTargetType.BALANCE,
        procedure_target_label="Cash and equivalents",
        expected_range_low=100000.0,
        expected_range_high=125000.0,
        precision_threshold_amount=5000.0,
        corroboration_basis_text="Budget review",
        corroboration_tags=["budget"],
        cpa_notes="Spike attributable to year-end customer pre-payment.",
    )
    mgr.update_expectation(user_id=user.id, expectation_id=exceeds.id, result_actual_value=140000.0)
    return AnalyticalExpectationMemoGenerator(db).generate_pdf(user.id, eng.id)


def _sum_schedule(request: Any) -> bytes:
    from sum_schedule_memo_generator import SumScheduleMemoGenerator
    from uncorrected_misstatements_manager import UncorrectedMisstatementsManager
    from uncorrected_misstatements_model import MisstatementClassification, MisstatementSourceType

    db, user, eng = _engagement(request)
    mgr = UncorrectedMisstatementsManager(db)
    for cls, amount in [
        (MisstatementClassification.FACTUAL, "-10000.00"),
        (MisstatementClassification.JUDGMENTAL, "-25000.00"),
        (MisstatementClassification.PROJECTED, "-90000.00"),
    ]:
        mgr.create_misstatement(
            user_id=user.id,
            engagement_id=eng.id,
            source_type=MisstatementSourceType.KNOWN_ERROR,
            source_reference=f"{cls.value} test",
            description=f"Test misstatement for {cls.value}",
            accounts_affected=[{"account": "Revenue", "debit_credit": "CR", "amount": 10000}],
            classification=cls,
            fs_impact_net_income=Decimal(amount),
            fs_impact_net_assets=Decimal("0.00"),
        )
    return SumScheduleMemoGenerator(db).generate_pdf(user.id, eng.id)


CASES: dict[str, Callable[[Any], bytes]] = {
    **{name.removeprefix("gen_"): _sample(name) for name in SAMPLE_GENERATORS},
    "engagement_dashboard": _engagement_dashboard,
    "analytical_expectations": _analytical_expectations,
    "sum_schedule": _sum_schedule,
}


# ---------------------------------------------------------------------------
# Snapshot plumbing
# ---------------------------------------------------------------------------


def _stable_hash(value: str) -> int:
    """``generate_testing_memo`` seeds its follow-up rotation with ``hash()``,
    which is salted per process. Pin it so the snapshot is reproducible."""
    return zlib.crc32(str(value).encode())


@pytest.fixture
def render(monkeypatch):
    import shared.memo_template as memo_template

    monkeypatch.setattr(memo_template, "hash", _stable_hash, raising=False)

    def _render(case: str, branding: str, request: Any) -> dict[str, Any]:
        with freeze_time(FROZEN_AT), apply_pdf_branding(BRANDINGS[branding]):
            pdf = CASES[case](request)
        return _measure(pdf)

    return _render


_recorded: dict[str, Any] = {}


def _load_snapshot() -> dict[str, Any]:
    if not SNAPSHOT_PATH.exists():
        return {}
    return json.loads(SNAPSHOT_PATH.read_text(encoding="utf-8"))


@pytest.fixture(scope="module")
def snapshot():
    data = _load_snapshot()
    yield data
    if UPDATE and _recorded:
        merged = {**data, **_recorded}
        SNAPSHOT_PATH.parent.mkdir(parents=True, exist_ok=True)
        SNAPSHOT_PATH.write_text(
            json.dumps(dict(sorted(merged.items())), indent=1, ensure_ascii=False) + "\n",
            encoding="utf-8",
        )


@pytest.mark.parametrize("branding", sorted(BRANDINGS))
@pytest.mark.parametrize("case", sorted(CASES))
def test_memo_renders_identically(case, branding, render, snapshot, request):
    key = f"{case}[{branding}]"
    actual = render(case, branding, request)
    if UPDATE:
        _recorded[key] = actual
        return

    assert key in snapshot, f"no recorded snapshot for {key}; re-record with MEMO_SNAPSHOT_UPDATE=1"
    expected = snapshot[key]
    assert actual["page_count"] == expected["page_count"], f"{key}: page count moved"
    for n, (got, want) in enumerate(zip(actual["pages"], expected["pages"], strict=True), start=1):
        for field in ("size", "text", "text_bbox", "runs", "content"):
            assert got[field] == want[field], f"{key} page {n}: {field} changed"


def test_snapshot_covers_every_case(snapshot):
    if UPDATE:
        pytest.skip("recording")
    expected = {f"{case}[{b}]" for case in CASES for b in BRANDINGS} | {"excel_workpaper"}
    assert set(snapshot) == expected


# ---------------------------------------------------------------------------
# Excel workpaper — item 9 hoists its per-cell fonts/fills to module constants
# ---------------------------------------------------------------------------


def _measure_xlsx(data: bytes) -> dict[str, Any]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data))
    sheets: dict[str, Any] = {}
    for ws in wb.worksheets:
        cells = []
        for row in ws.iter_rows():
            for c in row:
                if c.value is None and not c.has_style:
                    continue
                f, fill, b = c.font, c.fill, c.border
                cells.append(
                    [
                        c.coordinate,
                        repr(c.value),
                        c.style,
                        c.number_format,
                        [f.name, f.b, f.i, f.sz, str(getattr(f.color, "rgb", None))],
                        [fill.fill_type, str(getattr(fill.fgColor, "rgb", None))],
                        [
                            [getattr(s, "style", None), str(getattr(getattr(s, "color", None), "rgb", None))]
                            for s in (b.left, b.right, b.top, b.bottom)
                        ],
                        [c.alignment.horizontal, c.alignment.wrap_text],
                    ]
                )
        sheets[ws.title] = {
            "cells": _digest(cells),
            "cell_count": len(cells),
            "merged": sorted(str(r) for r in ws.merged_cells.ranges),
            "widths": {k: v.width for k, v in sorted(ws.column_dimensions.items())},
        }
    return {"sheet_order": wb.sheetnames, "sheets": sheets}


def _tb_audit_result(monkeypatch) -> dict[str, Any]:
    """The audit result the TB diagnostic sample renders, captured in flight."""
    import generate_sample_reports as samples
    import pdf
    import pdf_generator

    captured: dict[str, Any] = {}

    def capture(audit_result, *args, **kwargs):
        captured["audit_result"] = audit_result
        return b""

    # The sample imports it from ``pdf`` (Sprint 813) — patch the shim too.
    monkeypatch.setattr(pdf, "generate_audit_report", capture)
    monkeypatch.setattr(pdf_generator, "generate_audit_report", capture)
    monkeypatch.setattr(samples, "save_pdf", lambda *_a: None)
    samples.gen_tb_diagnostic()
    return captured["audit_result"]


def test_excel_workpaper_renders_identically(snapshot, monkeypatch):
    from excel_generator import generate_workpaper

    with freeze_time(FROZEN_AT):
        actual = _measure_xlsx(generate_workpaper(_tb_audit_result(monkeypatch), "meridian_tb.xlsx"))
    key = "excel_workpaper"
    if UPDATE:
        _recorded[key] = actual
        return
    assert actual == snapshot[key]
